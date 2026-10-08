"""生成后自检：启发式估算文本框是否溢出，溢出则逐步缩小字号并重存。

python-pptx 无法真实渲染，这里用「字数 × 字宽 vs 框宽高」做保守估算：
每段按字符数估算行数，累加行高，超过框高即判定溢出，然后对整框 run 字号逐级 -2pt 直到拟合。
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

_CHAR_W_FACTOR = 1.05   # 中文全角约 1 个字宽 ≈ 1×字号
_LINE_H_FACTOR = 1.35   # 行高 ≈ 1.35×字号
_MIN_SIZE = 10.0


def _est_height(tf, w_in: float) -> float:
    total = 0.0
    for para in tf.paragraphs:
        text = "".join(r.text for r in para.runs)
        if not text.strip():
            total += 0.18
            continue
        sizes = [r.font.size.pt for r in para.runs if r.font.size]
        size = max(sizes) if sizes else 18.0
        char_w = size / 72 * _CHAR_W_FACTOR
        per_line = max(int(w_in / char_w), 1)
        lines = (len(text) + per_line - 1) // per_line
        total += lines * (size / 72 * _LINE_H_FACTOR)
    return total


def _shrink(tf, step: float = 2.0) -> bool:
    changed = False
    for para in tf.paragraphs:
        for r in para.runs:
            if r.font.size:
                new = max(r.font.size.pt - step, _MIN_SIZE)
                if new < r.font.size.pt:
                    r.font.size = Pt(new)
                    changed = True
    return changed


def validate(path, out_path=None) -> list[str]:
    """检查并修复文字溢出，返回调整报告（字符串列表）。"""
    prs = Presentation(str(path))
    report: list[str] = []
    for idx, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            tf = shape.text_frame
            if not any(r.text.strip() for p in tf.paragraphs for r in p.runs):
                continue
            w_in = (shape.width or 0) / 914400
            h_in = (shape.height or 0) / 914400
            if w_in <= 0 or h_in <= 0:
                continue
            for _ in range(6):  # 最多缩 6 档
                if _est_height(tf, w_in) <= h_in * 1.02:
                    break
                if not _shrink(tf):
                    break
                report.append(f"第{idx}页 一个文本框字号已缩小（估算溢出）")
    out = out_path or path
    prs.save(str(out))
    return report
