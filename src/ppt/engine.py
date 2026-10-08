"""PPT 生成引擎（纯 python-pptx，不碰 LLM）。

两种生成方式：
- fill_template：直接套模板 —— 把大纲内容填进模板自带的「封面 + 内容页」。
- build_from_style：仿风格重新做 —— 抽取模板配色/字体，重建标准版式的新课件。

模板约定（直接套时）：第 1 页 = 封面（标题 + 副标题占位），第 2 页 = 内容页（标题 + 正文占位）。
"""
from __future__ import annotations

import copy
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_FILL
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from ..config import PPT_OUTPUT_DIR, TEMPLATES_DIR, TEMPLATE_SLOTS

# 内置默认风格（仿风格无模板时兜底）：墨绿封面 + 白底内容页，呼应系统整体配色。
DEFAULT_STYLE: dict = {
    "width": Inches(13.333),
    "height": Inches(7.5),
    "cover_bg": RGBColor(0x16, 0x28, 0x2B),
    "content_bg": RGBColor(0xFF, 0xFF, 0xFF),
    "accent": RGBColor(0x2F, 0xB5, 0x8F),
    "cover_title_font": "微软雅黑",
    "cover_title_size": Pt(40),
    "cover_title_color": RGBColor(0xFF, 0xFF, 0xFF),
    "subtitle_font": "微软雅黑",
    "subtitle_size": Pt(20),
    "subtitle_color": RGBColor(0xBF, 0xD3, 0xD0),
    "title_font": "微软雅黑",
    "title_size": Pt(30),
    "title_color": RGBColor(0x16, 0x28, 0x2B),
    "body_font": "微软雅黑",
    "body_size": Pt(20),
    "body_color": RGBColor(0x33, 0x33, 0x33),
}

_SLOTS_JSON = TEMPLATES_DIR / "slots.json"


# ---------------------------------------------------------------------------
# 槽位管理（5 个固定槽位）
# ---------------------------------------------------------------------------
def slot_path(slot: int) -> Path:
    return TEMPLATES_DIR / f"slot_{slot}.pptx"


def slot_names() -> dict[str, str]:
    try:
        return json.loads(_SLOTS_JSON.read_text(encoding="utf-8"))
    except Exception:
        return {}


def slot_filled(slot: int) -> bool:
    return slot_path(slot).exists()


def save_slot(slot: int, data: bytes, original_name: str) -> None:
    slot_path(slot).write_bytes(data)
    names = slot_names()
    names[str(slot)] = original_name
    _SLOTS_JSON.write_text(json.dumps(names, ensure_ascii=False, indent=2), encoding="utf-8")


# ---------------------------------------------------------------------------
# 直接套模板
# ---------------------------------------------------------------------------
def fill_template(template_path, outline: dict, out_path) -> None:
    prs = Presentation(str(template_path))
    if len(prs.slides) < 2:
        raise ValueError("模板至少需要 2 页（第 1 页封面、第 2 页内容页），否则请改用「仿风格」")

    cover = prs.slides[0]
    content_slide = prs.slides[1]

    _fill_text_by_idx(cover, 0, outline.get("title", ""))
    subtitle = (outline.get("subtitle") or "").strip()
    if subtitle:
        _fill_text_by_idx(cover, 1, subtitle)

    slides_data = outline.get("slides") or []
    n = max(len(slides_data), 1)
    # 先把空的内容页复制够，再逐页填，避免复制到已填内容。
    content_slides = [content_slide] + [_duplicate_slide(prs, content_slide) for _ in range(n - 1)]
    for slide, item in zip(content_slides, slides_data):
        _fill_text_by_idx(slide, 0, item.get("title", ""))
        _fill_bullets_by_idx(slide, 1, item.get("bullets") or [])

    prs.save(str(out_path))


def _duplicate_slide(prs, source):
    """复制一页（含形状文本；图片等需外部关系的内容页可能丢图，正文模板通常纯文字无影响）。"""
    new = prs.slides.add_slide(source.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    for shp in source.shapes:
        new.shapes._spTree.append(copy.deepcopy(shp._element))
    return new


# ---------------------------------------------------------------------------
# 仿风格重新做
# ---------------------------------------------------------------------------
def extract_style(template_path) -> dict:
    prs = Presentation(str(template_path))
    style = dict(DEFAULT_STYLE)
    style["width"] = prs.slide_width
    style["height"] = prs.slide_height

    slides = list(prs.slides)
    if not slides:
        return style

    cover = slides[0]
    cover_bg = _slide_bg_rgb(cover)
    if cover_bg is None:
        for shape in sorted(cover.shapes, key=lambda s: -_area(s)):
            c = _shape_fill_rgb(shape)
            if c is not None:
                cover_bg = c
                break
    if cover_bg is not None:
        style["cover_bg"] = cover_bg

    for shape in cover.shapes:
        c = _shape_fill_rgb(shape)
        if c is not None and c != cover_bg:
            style["accent"] = c
            break

    t = _placeholder_props(cover, 0) or _first_text_props(cover)
    if t:
        if t["name"]:
            style["cover_title_font"] = t["name"]
        if t["size"]:
            style["cover_title_size"] = t["size"]
        if t["color"]:
            style["cover_title_color"] = t["color"]
    sub = _placeholder_props(cover, 1)
    if sub:
        if sub["name"]:
            style["subtitle_font"] = sub["name"]
        if sub["size"]:
            style["subtitle_size"] = sub["size"]
        if sub["color"]:
            style["subtitle_color"] = sub["color"]

    if len(slides) >= 2:
        content = slides[1]
        cbg = _slide_bg_rgb(content)
        if cbg is not None:
            style["content_bg"] = cbg
        ct = _placeholder_props(content, 0)
        if ct:
            if ct["name"]:
                style["title_font"] = ct["name"]
            if ct["size"]:
                style["title_size"] = ct["size"]
            if ct["color"]:
                style["title_color"] = ct["color"]
        body = _placeholder_props(content, 1)
        if body:
            if body["name"]:
                style["body_font"] = body["name"]
            if body["size"]:
                style["body_size"] = body["size"]
            if body["color"]:
                style["body_color"] = body["color"]

    return style


def build_from_style(style: dict, outline: dict, out_path) -> None:
    prs = Presentation()
    prs.slide_width = style["width"]
    prs.slide_height = style["height"]
    blank = prs.slide_layouts[6]

    _build_cover(prs, blank, style, outline)
    for item in (outline.get("slides") or []):
        _build_content(prs, blank, style, item)

    prs.save(str(out_path))


def make_sample_template(path) -> None:
    """生成一个干净的示例模板（封面：标题+副标题；内容页：标题+正文），用于无模板时测试「直接套」。"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.slides.add_slide(prs.slide_layouts[0])  # Title Slide（标题 + 副标题）
    prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content（标题 + 正文）
    prs.save(str(path))


# ---------------------------------------------------------------------------
# 内部辅助：文本填充
# ---------------------------------------------------------------------------
def _fill_text_by_idx(slide, idx: int, text: str) -> bool:
    ph = _placeholder(slide, idx)
    if ph is not None and ph.has_text_frame:
        ph.text_frame.text = text
        return True
    shapes = _text_shapes(slide)
    if len(shapes) > idx:
        shapes[idx].text_frame.text = text
        return True
    if shapes:
        shapes[0].text_frame.text = text
        return True
    return False


def _fill_bullets_by_idx(slide, idx: int, bullets: list[str]) -> bool:
    ph = _placeholder(slide, idx)
    if ph is not None and ph.has_text_frame:
        _set_bullets(ph.text_frame, bullets)
        return True
    shapes = _text_shapes(slide)
    if len(shapes) > idx:
        _set_bullets(shapes[idx].text_frame, bullets)
        return True
    return False


def _placeholder(slide, idx: int):
    try:
        return slide.placeholders[idx]
    except Exception:
        return None


def _text_shapes(slide) -> list:
    shapes = [s for s in slide.shapes if s.has_text_frame]
    shapes.sort(key=lambda s: (s.top or 0, s.left or 0))
    return shapes


def _set_bullets(tf, bullets: list[str]) -> None:
    tf.clear()
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b


# ---------------------------------------------------------------------------
# 内部辅助：样式抽取
# ---------------------------------------------------------------------------
def _area(shape) -> int:
    return int((shape.width or 0) * (shape.height or 0))


def _shape_fill_rgb(shape):
    try:
        if shape.fill.type == MSO_FILL.SOLID:
            return shape.fill.fore_color.rgb
    except Exception:
        pass
    return None


def _slide_bg_rgb(slide):
    try:
        if slide.background.fill.type == MSO_FILL.SOLID:
            return slide.background.fill.fore_color.rgb
    except Exception:
        pass
    return None


def _run_props(run) -> dict:
    f = run.font
    color = None
    try:
        color = f.color.rgb
    except Exception:
        color = None
    return {"name": f.name, "size": f.size, "color": color, "bold": bool(f.bold)}


def _first_text_props(slide):
    for shape in _text_shapes(slide):
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                if run.text.strip():
                    return _run_props(run)
    return None


def _placeholder_props(slide, idx: int):
    ph = _placeholder(slide, idx)
    if ph is None or not ph.has_text_frame:
        return None
    for para in ph.text_frame.paragraphs:
        for run in para.runs:
            if run.text.strip():
                return _run_props(run)
    return None


# ---------------------------------------------------------------------------
# 内部辅助：从零建页
# ---------------------------------------------------------------------------
def _set_run_font(run, name, size=None, color=None, bold=None) -> None:
    if name:
        run.font.name = name
        # 同步设置东亚字体，否则中文不会用指定字体
        rPr = run._r.get_or_add_rPr()
        ea = rPr.find(qn("a:ea"))
        if ea is None:
            ea = rPr.makeelement(qn("a:ea"), {})
            rPr.append(ea)
        ea.set("typeface", name)
    if size is not None:
        run.font.size = size
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.font.bold = bold


def _add_rect(slide, left, top, width, height, fill_rgb):
    shp = slide.shapes.add_shape(1, left, top, width, height)  # MSO_SHAPE.RECTANGLE == 1
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill_rgb
    shp.line.fill.background()
    try:
        shp.shadow.inherit = False
    except Exception:
        pass
    return shp


def _add_textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    box.text_frame.word_wrap = True
    return box


def _set_para(tf, text: str, font, size, color, bold=False, align=PP_ALIGN.LEFT) -> None:
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    if p.runs:
        _set_run_font(p.runs[0], font, size, color, bold)


def _set_bullet_paras(tf, bullets: list[str], font, size, color) -> None:
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "• " + b
        p.space_after = Pt(8)
        for run in p.runs:
            _set_run_font(run, font, size, color)


def _build_cover(prs, layout, style, outline) -> None:
    slide = prs.slides.add_slide(layout)
    W, H = prs.slide_width, prs.slide_height
    _add_rect(slide, 0, 0, W, H, style["cover_bg"])
    _add_rect(slide, 0, 0, Inches(0.35), H, style["accent"])

    tb = _add_textbox(slide, Inches(1.0), Inches(2.5), W - Inches(2.0), Inches(1.7))
    _set_para(tb.text_frame, outline.get("title", ""), style["cover_title_font"],
              style["cover_title_size"], style["cover_title_color"], bold=True)

    subtitle = (outline.get("subtitle") or "").strip()
    if subtitle:
        tb2 = _add_textbox(slide, Inches(1.0), Inches(4.4), W - Inches(2.0), Inches(0.8))
        _set_para(tb2.text_frame, subtitle, style["subtitle_font"],
                  style["subtitle_size"], style["subtitle_color"])


def _build_content(prs, layout, style, item) -> None:
    slide = prs.slides.add_slide(layout)
    W, H = prs.slide_width, prs.slide_height
    _add_rect(slide, 0, 0, W, H, style["content_bg"])
    _add_rect(slide, 0, 0, Inches(0.22), H, style["accent"])

    tb = _add_textbox(slide, Inches(0.7), Inches(0.5), W - Inches(1.4), Inches(1.0))
    _set_para(tb.text_frame, item.get("title", ""), style["title_font"],
              style["title_size"], style["title_color"], bold=True)

    tb2 = _add_textbox(slide, Inches(0.9), Inches(1.8), W - Inches(1.8), H - Inches(2.4))
    _set_bullet_paras(tb2.text_frame, item.get("bullets") or [],
                      style["body_font"], style["body_size"], style["body_color"])
