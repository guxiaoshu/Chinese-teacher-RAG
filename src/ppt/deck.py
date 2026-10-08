"""通用设计引擎：按 Style 渲染多版式课件（封面/分节/内容/金句/对比/流程/封底）。

- 风格由 styles.Style 驱动（配色/字体/装饰/配图后缀）
- 每页 layout 决定版式；LLM 输出 layout 字段，缺省按 content
- 配图：有 image_provider 且该页有 image_prompt → 真图满铺 + 半透明覆盖层；
  否则回退到纯矢量装饰（竹影/石潭/游鱼/几何）
- 生成后调用 validate 做文字溢出自检

入口：build_deck(outline, style, out_path, image_provider=None) -> list[str]（自检报告）
"""
from __future__ import annotations

import os
import tempfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .imagegen import ImageProvider
from .styles import Style
from .validate import validate

SW, SH = 13.333, 7.5  # 16:9 英寸
_CN = "零一二三四五六七八九"


def _cn_num(n: int) -> str:
    if n <= 0:
        return ""
    if n <= 9:
        return _CN[n]
    if n == 10:
        return "十"
    if n < 20:
        return "十" + _CN[n - 10]
    if n < 100:
        t, o = n // 10, n % 10
        return _CN[t] + "十" + (_CN[o] if o else "")
    return str(n)


# ---------------------------------------------------------------------------
# 基础工具（颜色/字体显式传入，风格无关）
# ---------------------------------------------------------------------------
def _set_alpha(shape, pct):
    val = str(int(round(pct * 1000)))
    spPr = shape._element.spPr
    solidFill = spPr.find(qn("a:solidFill"))
    if solidFill is None:
        return
    srgb = solidFill.find(qn("a:srgbClr"))
    if srgb is None:
        return
    for e in srgb.findall(qn("a:alpha")):
        srgb.remove(e)
    a = srgb.makeelement(qn("a:alpha"), {"val": val})
    srgb.append(a)


def _set_run_font(run, name, size, color, bold=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        latin = rPr.find(qn("a:latin"))
        if latin is not None:
            latin.addnext(ea)
        else:
            rPr.append(ea)
    ea.set("typeface", name)


def _rect(slide, x, y, w, h, fill, alpha=None, rounded=False, radius=0.1):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if alpha is not None:
        _set_alpha(shp, alpha)
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _oval(slide, x, y, w, h, fill, alpha=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if alpha is not None:
        _set_alpha(shp, alpha)
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _add_tb(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tb, tf


def _add_bg(slide, style):
    _rect(slide, 0, 0, SW, SH, style.bg)


def _add_seal(slide, chars, x, y, size=0.5, fill=None):
    fill = fill or RGBColor(0xC0, 0x39, 0x2B)
    shp = _rect(slide, x, y, size, size, fill, rounded=True, radius=0.16)
    tf = shp.text_frame
    tf.word_wrap = False
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    n = max(len(chars), 1)
    fsz = size * 72 * (0.42 if n == 1 else 0.36)
    for i, ch in enumerate(chars):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.line_spacing = 0.9
        r = p.add_run()
        r.text = ch
        _set_run_font(r, "楷体", fsz, RGBColor(0xFF, 0xFF, 0xFF), False)
    return shp


def _vtext(slide, chars, x, y, size, color, bold=False, font="楷体", ls=1.0):
    for i, ch in enumerate(chars):
        tb, tf = _add_tb(slide, x, y + i * (size / 72) * ls, size / 72 * 1.4 + 0.2, size / 72 * ls + 0.2)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = ch
        _set_run_font(r, font, size, color, bold)


def _htext(slide, text, x, y, w, size, color, bold=False, font="宋体", align=PP_ALIGN.LEFT, ls=1.15):
    tb, tf = _add_tb(slide, x, y, w, size / 72 * 2.4 + 0.5)
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = ls
    r = p.add_run()
    r.text = text
    _set_run_font(r, font, size, color, bold)
    return tb


def _bullets(slide, items, x, y, w, size=18, gap=0.62, marker="·", marker_color=None, color=None):
    tb, tf = _add_tb(slide, x, y, w, len(items) * gap + 0.4)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap * 72 * 0.18)
        p.line_spacing = 1.12
        rm = p.add_run()
        rm.text = marker + " "
        _set_run_font(rm, "楷体", size, marker_color, True)
        r = p.add_run()
        r.text = it
        _set_run_font(r, "宋体", size, color, False)
    return tb


def _vt_size(n):
    return 30 if n <= 4 else (24 if n <= 6 else (20 if n <= 8 else 18))


# ---------------------------------------------------------------------------
# 矢量装饰（无真图时的兜底）
# ---------------------------------------------------------------------------
def _add_bamboo(slide, x, top, hgt, style, count=3, gap=0.24):
    for i in range(count):
        c = style.primary if i % 2 == 0 else style.secondary
        _rect(slide, x + i * gap, top, 0.035, hgt, c, alpha=55)


def _add_pond_art(slide, has_fish, style):
    _oval(slide, 8.5, 4.6, 3.7, 1.7, style.secondary, 28)
    _oval(slide, 9.25, 5.15, 2.2, 0.85, style.primary, 18)
    _oval(slide, 8.15, 4.3, 0.5, 0.34, style.gold, 55)
    if has_fish:
        for i, (fx, fy, l) in enumerate([(9.2, 5.5, 0.5), (9.85, 5.28, 0.44), (10.4, 5.52, 0.4)]):
            _oval(slide, fx, fy, l, l * 0.32, style.primary if i % 2 == 0 else style.secondary, 75)


def _decorate(slide, style, has_fish=False):
    if style.key == "minimal":
        _oval(slide, 9.8, 1.0, 2.3, 2.3, style.primary, 12)
        _rect(slide, 12.55, 1.3, 0.05, 4.6, style.gold)
    else:
        _add_bamboo(slide, 12.75, 0.9, 5.6, style)
        _add_pond_art(slide, has_fish, style)


# ---------------------------------------------------------------------------
# 真图
# ---------------------------------------------------------------------------
def _cover_prompt(outline, style):
    parts = [x for x in [outline.get("title", ""), outline.get("subtitle", "")] if x]
    return "，".join(parts) + "，" + style.image_suffix


def _slide_prompt(item, style):
    p = (item.get("image_prompt") or "").strip()
    return f"{p}，{style.image_suffix}" if p else ""


def _gen_images(outline, style, provider, max_images=6):
    if provider is None:
        return {}
    prompts = [_cover_prompt(outline, style)]
    for it in outline.get("slides") or []:
        p = _slide_prompt(it, style)
        if p:
            prompts.append(p)
    seen = []
    for p in prompts:
        if p and p not in seen:
            seen.append(p)
    seen = seen[:max_images]

    def work(p):
        try:
            return p, provider.generate(p)
        except Exception:
            return p, None

    images = {}
    with ThreadPoolExecutor(max_workers=min(4, max(len(seen), 1))) as ex:
        for p, data in ex.map(work, seen):
            if data:
                images[p] = data
    return images


def _set_bg_image(slide, data, style, overlay_alpha=55):
    tmp = None
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
        f.write(data)
        tmp = f.name
    try:
        slide.shapes.add_picture(tmp, 0, 0, Inches(SW), Inches(SH))
    finally:
        if tmp:
            os.unlink(tmp)
    _rect(slide, 0, 0, SW, SH, style.bg, alpha=overlay_alpha)


# ---------------------------------------------------------------------------
# 版式 builder
# ---------------------------------------------------------------------------
def _title(slide, style, title):
    """标题：≤8 字竖排，否则横排。返回内容区起点 (x0, y0)。"""
    n = len(title)
    if n <= 8:
        _vtext(slide, title, 0.55, 0.62, _vt_size(n), style.ink, True, style.title_font)
        _rect(slide, 1.55, 0.72, 0.03, min(n, 8) * 0.46, style.gold)
        return 2.0, 1.0
    _htext(slide, title, 0.7, 0.42, SW - 1.4, 24, style.ink, True, style.title_font)
    _rect(slide, 0.7, 1.02, 1.2, 0.03, style.gold)
    return 0.7, 1.4


def _page_seal(slide, style, pageno):
    _add_seal(slide, _cn_num(pageno), 0.55, 6.75, 0.5, style.accent)
    _oval(slide, 12.55, 0.55, 0.09, 0.09, style.gold, 90)


def _set_note(slide, item):
    note = (item.get("note") or "").strip()
    if note:
        slide.notes_slide.notes_text_frame.text = note


def _build_cover(prs, style, outline, images):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    img = images.get(_cover_prompt(outline, style))
    if img:
        _set_bg_image(slide, img, style, overlay_alpha=38)
    else:
        _add_bg(slide, style)
        _decorate(slide, style)

    title = (outline.get("title") or "").strip() or "课件"
    subtitle = (outline.get("subtitle") or "").strip()

    _add_seal(slide, title[:2] if len(title) >= 2 else "课件", 12.05, 0.55, 0.72, style.accent)
    tt = title[:8]
    _vtext(slide, tt, 1.15, 1.5, _vt_size(len(tt)) + 24, style.ink, True, style.title_font)
    _rect(slide, 3.15, 1.7, 0.03, 3.1, style.gold)
    if subtitle:
        _htext(slide, subtitle, 1.15, 6.35, 6.5, 16, style.primary, font=style.title_font)


def _build_section(prs, style, item, pageno, images):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    img = images.get(_slide_prompt(item, style))
    if img:
        _set_bg_image(slide, img, style, overlay_alpha=52)
    else:
        _add_bg(slide, style)
    _page_seal(slide, style, pageno)
    title = (item.get("title") or "").strip() or "分节"
    _rect(slide, (SW - 1.2) / 2, 2.3, 1.2, 0.03, style.gold)
    _htext(slide, title, 1.5, 2.7, SW - 3.0, 44 if len(title) <= 6 else 34,
           style.ink, True, style.title_font, align=PP_ALIGN.CENTER, ls=1.2)
    _set_note(slide, item)


def _build_content(prs, style, item, pageno, images):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = (item.get("title") or "").strip() or "内容"
    quote = (item.get("quote") or "").strip()
    bullets = [str(b).strip() for b in (item.get("bullets") or []) if str(b).strip()]
    hay = title + quote + "".join(bullets)

    img = images.get(_slide_prompt(item, style))
    if img:
        _set_bg_image(slide, img, style, overlay_alpha=58)
    else:
        _add_bg(slide, style)
        _decorate(slide, style, has_fish=("鱼" in hay))

    _page_seal(slide, style, pageno)
    x0, y0 = _title(slide, style, title)

    if quote:
        _rect(slide, x0, y0, 0.06, 0.5, style.accent)
        _htext(slide, quote, x0 + 0.25, y0, 6.0, 20, style.primary, True, style.title_font)
        y = y0 + 1.0
    else:
        y = y0 + 0.2
    if bullets:
        _bullets(slide, bullets, x0, y, 6.2, size=18, gap=0.62,
                 marker="·", marker_color=style.accent, color=style.ink)
    _set_note(slide, item)


def _build_quote(prs, style, item, pageno, images):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    img = images.get(_slide_prompt(item, style))
    if img:
        _set_bg_image(slide, img, style, overlay_alpha=46)
    else:
        _add_bg(slide, style)
        _decorate(slide, style)
    _page_seal(slide, style, pageno)
    quote = (item.get("quote") or item.get("title") or "").strip() or "金句"
    n = len(quote)
    size = 32 if n <= 12 else (28 if n <= 20 else 22)
    _rect(slide, (SW - 1.2) / 2, 2.3, 1.2, 0.03, style.gold)
    _htext(slide, quote, 1.6, 2.7, SW - 3.2, size, style.ink, True, style.title_font,
           align=PP_ALIGN.CENTER, ls=1.35)
    attribution = (item.get("title") or "").strip()
    if attribution and attribution != quote:
        _htext(slide, attribution, 1.6, 5.4, SW - 3.2, 15, style.primary,
               font=style.title_font, align=PP_ALIGN.CENTER)
    _set_note(slide, item)


def _build_compare(prs, style, item, pageno, images):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    img = images.get(_slide_prompt(item, style))
    if img:
        _set_bg_image(slide, img, style, overlay_alpha=66)
    else:
        _add_bg(slide, style)
    _page_seal(slide, style, pageno)
    title = (item.get("title") or "").strip() or "对比"
    _htext(slide, title, 0.7, 0.4, SW - 1.4, 24, style.ink, True, style.title_font)
    _rect(slide, 0.7, 1.0, 1.2, 0.03, style.gold)

    cmp = item.get("compare") or {}
    left = cmp.get("left") or {}
    right = cmp.get("right") or {}

    cw, cy, ch = 5.4, 1.35, 5.0
    for cx, side in ((0.9, left), (7.0, right)):
        _rect(slide, cx, cy, cw, ch, style.secondary, 16, rounded=True, radius=0.04)
        stitle = (side.get("title") or "").strip()
        sbullets = [str(b).strip() for b in (side.get("bullets") or []) if str(b).strip()]
        if stitle:
            _htext(slide, stitle, cx + 0.3, cy + 0.25, cw - 0.6, 19, style.ink, True, style.title_font)
            by = cy + 1.05
        else:
            by = cy + 0.45
        if sbullets:
            _bullets(slide, sbullets, cx + 0.3, by, cw - 0.6, size=15, gap=0.55,
                     marker="·", marker_color=style.accent, color=style.ink)
    _set_note(slide, item)


def _build_flow(prs, style, item, pageno, images):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    img = images.get(_slide_prompt(item, style))
    if img:
        _set_bg_image(slide, img, style, overlay_alpha=62)
    else:
        _add_bg(slide, style)
    _page_seal(slide, style, pageno)
    title = (item.get("title") or "").strip() or "流程"
    _htext(slide, title, 0.7, 0.4, SW - 1.4, 24, style.ink, True, style.title_font)
    _rect(slide, 0.7, 1.0, 1.2, 0.03, style.gold)

    steps = [str(s).strip() for s in (item.get("steps") or []) if str(s).strip()]
    if steps:
        n = len(steps)
        node_w, node_h, gap = 2.05, 1.7, 0.5
        total = n * node_w + (n - 1) * gap
        x_start = (SW - total) / 2
        y = 2.85
        for i, st in enumerate(steps):
            nx = x_start + i * (node_w + gap)
            _rect(slide, nx, y, node_w, node_h, style.secondary, 18, rounded=True, radius=0.12)
            _rect(slide, nx, y, node_w, 0.06, style.accent)
            tb, tf = _add_tb(slide, nx + 0.12, y + 0.35, node_w - 0.24, node_h - 0.6, anchor=MSO_ANCHOR.MIDDLE)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.line_spacing = 1.05
            r = p.add_run()
            r.text = st
            _set_run_font(r, style.body_font, 14, style.ink, False)
            if i < n - 1:
                _rect(slide, nx + node_w, y + node_h / 2 - 0.015, gap - 0.2, 0.03, style.gold)
                _oval(slide, nx + node_w + gap - 0.28, y + node_h / 2 - 0.07, 0.14, 0.14, style.gold)
    _set_note(slide, item)


def _build_closing(prs, style, item, pageno, images):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    img = images.get(_slide_prompt(item, style))
    if img:
        _set_bg_image(slide, img, style, overlay_alpha=44)
    else:
        _add_bg(slide, style)
        _decorate(slide, style)
    quote = (item.get("quote") or item.get("title") or "").strip() or "景中有情"
    _htext(slide, quote, 1.6, 2.9, SW - 3.2, 32, style.ink, True, style.title_font,
           align=PP_ALIGN.CENTER, ls=1.3)
    title = (item.get("title") or "").strip()
    if title and title != quote:
        _htext(slide, title, 1.6, 5.4, SW - 3.2, 15, style.primary, font=style.title_font,
               align=PP_ALIGN.CENTER)
    _add_seal(slide, _cn_num(pageno), 0.55, 6.75, 0.5, style.accent)
    _set_note(slide, item)


_LAYOUTS = {
    "section": _build_section,
    "content": _build_content,
    "quote": _build_quote,
    "compare": _build_compare,
    "flow": _build_flow,
    "closing": _build_closing,
}


# ---------------------------------------------------------------------------
# 入口
# ---------------------------------------------------------------------------
def build_deck(outline: dict, style: Style, out_path, image_provider: ImageProvider | None = None,
               max_images: int = 6) -> list[str]:
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)

    images = _gen_images(outline, style, image_provider, max_images)
    _build_cover(prs, style, outline, images)

    slides = [s for s in (outline.get("slides") or []) if isinstance(s, dict)]
    for i, item in enumerate(slides, 2):  # 封面第 1 页，内容页从 2 起
        layout = (item.get("layout") or "content").strip().lower()
        builder = _LAYOUTS.get(layout, _build_content)
        builder(prs, style, item, i, images)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return validate(out_path)
